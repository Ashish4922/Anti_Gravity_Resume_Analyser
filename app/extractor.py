"""Document text and metadata extractor supporting PDF, DOCX, and plain text.
"""

import io
import re
from typing import Dict, Any
import pypdf
import docx


class DocumentExtractor:
    """Extracts text, metadata, and structural cues from uploaded files."""

    @staticmethod
    def extract_from_bytes(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        """Extract text and metadata based on file extension."""
        lower_name = filename.lower()

        if lower_name.endswith(".pdf"):
            return DocumentExtractor._extract_pdf(file_bytes)
        elif lower_name.endswith(".docx"):
            return DocumentExtractor._extract_docx(file_bytes)
        elif lower_name.endswith(".txt") or lower_name.endswith(".md"):
            return DocumentExtractor._extract_text(file_bytes)
        else:
            raise ValueError(f"Unsupported file format: {filename}. Please upload PDF, DOCX, or TXT.")

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> Dict[str, Any]:
        """Extract text from PDF using pypdf."""
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        page_count = len(reader.pages)
        pages_text = []

        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            pages_text.append(page_text.strip())

        full_text = "\n\n".join(pages_text).strip()
        is_scanned = len(full_text.strip()) < 50 and page_count > 0

        return {
            "format": "PDF",
            "page_count": page_count,
            "text": full_text,
            "is_scanned": is_scanned,
            "char_count": len(full_text),
            "word_count": len(full_text.split())
        }

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> Dict[str, Any]:
        """Extract text from DOCX document including paragraphs and tables."""
        doc = docx.Document(io.BytesIO(file_bytes))
        lines = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                lines.append(text)

        # Extract table content
        table_lines = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    table_lines.append(row_text)

        all_text = "\n".join(lines)
        if table_lines:
            all_text += "\n\n" + "\n".join(table_lines)

        word_count = len(all_text.split())
        # Estimate page count for DOCX (~450 words per page)
        estimated_pages = max(1, round(word_count / 450)) if word_count > 0 else 1

        return {
            "format": "DOCX",
            "page_count": estimated_pages,
            "text": all_text.strip(),
            "is_scanned": False,
            "char_count": len(all_text),
            "word_count": word_count
        }

    @staticmethod
    def _extract_text(file_bytes: bytes) -> Dict[str, Any]:
        """Extract text from plain text file with encoding fallbacks."""
        text = ""
        for encoding in ["utf-8", "utf-16", "cp1252", "latin-1"]:
            try:
                text = file_bytes.decode(encoding)
                break
            except (UnicodeDecodeError, LookupError):
                continue

        word_count = len(text.split())
        estimated_pages = max(1, round(word_count / 450)) if word_count > 0 else 1

        return {
            "format": "TXT",
            "page_count": estimated_pages,
            "text": text.strip(),
            "is_scanned": False,
            "char_count": len(text),
            "word_count": word_count
        }
